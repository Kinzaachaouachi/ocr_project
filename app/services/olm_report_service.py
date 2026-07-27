import io
import csv
from datetime import datetime
from typing import Optional

from sqlalchemy import func
from sqlalchemy.orm import Session

from ..models.ocr_history import OCRHistory
from ..utils.olm_benchmark_matrix import (
    OLM_BENCHMARK_MATRIX,
    BENCHMARK_METADATA,
    SCORING_METHODOLOGY,
    OLMOCR_CHANGELOG,
    CRITERIA_DESCRIPTIONS,
    CRITERIA_LABELS,
    CRITERIA_ICONS,
    get_top_models,
    get_model_strengths,
    get_model_weaknesses,
    get_performance_level,
    get_category_leaders,
    get_model_comparison_table,
)


def generate_olm_report_data(db: Session, user_id: Optional[int] = None) -> dict:
    """
    Generate comprehensive OLM Benchmark report data.

    Args:
        db: Database session
        user_id: Optional user ID to include user-specific stats

    Returns:
        Dictionary with all report data
    """

    benchmark_matrix = get_model_comparison_table()

    top_by_category = get_category_leaders()

    user_stats = None
    if user_id:
        total_extractions = (
            db.query(OCRHistory).filter(OCRHistory.user_id == user_id).count()
        )

        successful = db.query(OCRHistory).filter(
            OCRHistory.user_id == user_id, OCRHistory.status == "success"
        )

        success_count = successful.count()

        metrics = successful.with_entities(
            func.avg(OCRHistory.precision_score).label("avg_precision"),
            func.avg(OCRHistory.ocr_time_s).label("avg_exec_time"),
            func.avg(OCRHistory.robustness).label("avg_robustness"),
            func.avg(OCRHistory.global_score).label("avg_global_score"),
        ).first()

        model_usage = (
            successful.with_entities(
                OCRHistory.model_name, func.count(OCRHistory.id).label("count")
            )
            .group_by(OCRHistory.model_name)
            .all()
        )

        user_stats = {
            "total_extractions": total_extractions,
            "success_count": success_count,
            "success_rate": (
                round((success_count / total_extractions * 100), 2)
                if total_extractions > 0
                else 0
            ),
            "avg_precision": round(float(metrics.avg_precision or 0), 2),
            "avg_execution_time": round(float(metrics.avg_exec_time or 0), 2),
            "avg_robustness": round(float(metrics.avg_robustness or 0), 2),
            "avg_global_score": round(float(metrics.avg_global_score or 0), 2),
            "model_usage": [
                {"model": m.model_name, "count": m.count} for m in model_usage
            ],
        }

    return {
        "generated_at": datetime.now().isoformat(),
        "benchmark_info": {
            "name": BENCHMARK_METADATA["name"],
            "version": BENCHMARK_METADATA["version"],
            "description": BENCHMARK_METADATA["description"],
            "source": BENCHMARK_METADATA["source"],
            "github_url": BENCHMARK_METADATA["github_url"],
            "huggingface_url": BENCHMARK_METADATA["huggingface_url"],
            "demo_url": BENCHMARK_METADATA["demo_url"],
            "arxiv_paper_v1": BENCHMARK_METADATA["arxiv_paper_v1"],
            "arxiv_paper_v2": BENCHMARK_METADATA["arxiv_paper_v2"],
            "license": BENCHMARK_METADATA["license"],
            "test_categories": BENCHMARK_METADATA["test_categories"],
            "total_documents": BENCHMARK_METADATA["total_documents"],
            "total_tests": BENCHMARK_METADATA["total_tests"],
            "last_updated": BENCHMARK_METADATA["last_updated"],
        },
        "scoring_methodology": SCORING_METHODOLOGY,
        "changelog": OLMOCR_CHANGELOG,
        "criteria": {
            "labels": CRITERIA_LABELS,
            "descriptions": CRITERIA_DESCRIPTIONS,
            "icons": CRITERIA_ICONS,
        },
        "benchmark_matrix": benchmark_matrix,
        "top_by_category": top_by_category,
        "user_stats": user_stats,
    }


def generate_html_report(report_data: dict) -> str:
    """Generate comprehensive HTML report from report data."""

    generated = datetime.fromisoformat(report_data["generated_at"]).strftime(
        "%d/%m/%Y à %H:%M"
    )

    html = f"""<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Rapport OLM Benchmark Global — olmOCR-Bench</title>
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&display=swap');
        * {{ margin: 0; padding: 0; box-sizing: border-box; }}
        body {{ font-family: 'Inter', 'Segoe UI', sans-serif; background: #f0f2f5; padding: 20px; color: #333; line-height: 1.6; }}
        .container {{ max-width: 1200px; margin: 0 auto; background: white; border-radius: 16px; box-shadow: 0 4px 24px rgba(0,0,0,0.08); overflow: hidden; }}

        /* Header */
        .header {{ background: linear-gradient(135deg, #1e3a8a 0%, #3b82f6 50%, #6366f1 100%); color: white; padding: 40px; text-align: center; position: relative; }}
        .header::after {{ content: ''; position: absolute; bottom: 0; left: 0; right: 0; height: 4px; background: linear-gradient(90deg, #fbbf24, #f59e0b, #d97706); }}
        .header h1 {{ font-size: 2.4em; font-weight: 800; margin-bottom: 8px; letter-spacing: -0.5px; }}
        .header .subtitle {{ font-size: 1.15em; opacity: 0.9; font-weight: 400; }}
        .header .date {{ font-size: 0.9em; opacity: 0.7; margin-top: 12px; }}
        .header .badge {{ display: inline-block; background: rgba(255,255,255,0.2); padding: 4px 14px; border-radius: 20px; font-size: 0.85em; margin-top: 8px; backdrop-filter: blur(4px); }}

        /* Content sections */
        .content {{ padding: 30px 40px; }}
        .section {{ margin: 35px 0; }}
        .section h2 {{ color: #1e3a8a; font-size: 1.6em; font-weight: 700; border-bottom: 3px solid #e2e8f0; padding-bottom: 12px; margin-bottom: 20px; }}
        .section h3 {{ color: #334155; font-size: 1.2em; font-weight: 600; margin: 18px 0 10px; }}

        /* Info box */
        .info-box {{ background: linear-gradient(135deg, #eff6ff, #f0f4ff); border-left: 4px solid #2563eb; padding: 24px; margin: 20px 0; border-radius: 0 12px 12px 0; }}
        .info-box h3 {{ color: #1e3a8a; margin-bottom: 12px; font-size: 1.1em; }}
        .info-box p {{ margin: 6px 0; color: #475569; }}
        .info-box a {{ color: #2563eb; text-decoration: none; font-weight: 600; }}
        .info-box a:hover {{ text-decoration: underline; }}

        /* Methodology box */
        .methodology {{ background: #fefce8; border-left: 4px solid #eab308; padding: 20px; border-radius: 0 12px 12px 0; margin: 20px 0; }}
        .methodology h3 {{ color: #854d0e; margin-bottom: 10px; }}
        .methodology p {{ color: #713f12; margin: 6px 0; font-size: 0.95em; }}

        /* Performance levels */
        .perf-grid {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 10px; margin: 15px 0; }}
        .perf-item {{ padding: 10px 14px; border-radius: 8px; font-size: 0.85em; font-weight: 500; text-align: center; }}
        .perf-excellent {{ background: #dcfce7; color: #166534; border: 1px solid #bbf7d0; }}
        .perf-good {{ background: #dbeafe; color: #1e40af; border: 1px solid #bfdbfe; }}
        .perf-acceptable {{ background: #fef3c7; color: #92400e; border: 1px solid #fde68a; }}
        .perf-poor {{ background: #fee2e2; color: #991b1b; border: 1px solid #fecaca; }}

        /* Matrix table */
        .matrix-table {{ width: 100%; border-collapse: collapse; margin: 20px 0; font-size: 0.92em; }}
        .matrix-table th {{ background: linear-gradient(135deg, #1e3a8a, #2563eb); color: white; padding: 14px 10px; text-align: center; font-weight: 600; font-size: 0.88em; text-transform: uppercase; letter-spacing: 0.5px; }}
        .matrix-table td {{ border: 1px solid #e2e8f0; padding: 12px 10px; text-align: center; }}
        .matrix-table tr:nth-child(even) {{ background: #f8fafc; }}
        .matrix-table tr:hover {{ background: #eff6ff; }}
        .matrix-table .model-name {{ text-align: left; font-weight: 600; color: #1e293b; }}

        /* Rank badges */
        .rank-badge {{ display: inline-block; padding: 4px 12px; border-radius: 20px; font-weight: 700; font-size: 0.85em; }}
        .rank-1 {{ background: linear-gradient(135deg, #fbbf24, #f59e0b); color: #78350f; box-shadow: 0 2px 8px rgba(251,191,36,0.3); }}
        .rank-2 {{ background: linear-gradient(135deg, #d1d5db, #9ca3af); color: #374151; }}
        .rank-3 {{ background: linear-gradient(135deg, #d97706, #b45309); color: white; }}
        .rank-other {{ background: #e2e8f0; color: #64748b; }}

        /* Score coloring */
        .score-excellent {{ color: #059669; font-weight: 700; }}
        .score-high {{ color: #2563eb; font-weight: 600; }}
        .score-medium {{ color: #d97706; font-weight: 500; }}
        .score-low {{ color: #dc2626; font-weight: 600; }}

        /* Criteria grid */
        .criteria-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(280px, 1fr)); gap: 14px; margin: 20px 0; }}
        .criteria-card {{ background: #f8fafc; padding: 16px; border-radius: 10px; border-left: 4px solid #3b82f6; transition: transform 0.2s; }}
        .criteria-card:hover {{ transform: translateY(-2px); box-shadow: 0 4px 12px rgba(0,0,0,0.06); }}
        .criteria-card h4 {{ color: #1e3a8a; margin-bottom: 6px; font-size: 1em; }}
        .criteria-card .leader {{ font-size: 0.85em; color: #059669; margin-top: 8px; font-weight: 500; }}
        .criteria-card p {{ color: #64748b; font-size: 0.88em; }}

        /* Model detail cards */
        .model-card {{ margin: 18px 0; padding: 20px; background: #f8fafc; border-radius: 12px; border: 1px solid #e2e8f0; }}
        .model-card h3 {{ color: #1e3a8a; font-size: 1.1em; margin-bottom: 4px; }}
        .model-card .meta {{ font-size: 0.85em; color: #64748b; margin-bottom: 12px; }}
        .model-card ul {{ margin-left: 20px; }}
        .model-card li {{ margin: 4px 0; font-size: 0.92em; }}

        /* Stats grid */
        .stats-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(200px, 1fr)); gap: 16px; margin: 20px 0; }}
        .stat-card {{ background: linear-gradient(135deg, #667eea 0%, #764ba2 100%); color: white; padding: 20px; border-radius: 12px; }}
        .stat-card h4 {{ font-size: 0.88em; opacity: 0.9; margin-bottom: 8px; font-weight: 500; }}
        .stat-card .value {{ font-size: 2em; font-weight: 800; }}

        /* Changelog */
        .changelog {{ margin: 20px 0; }}
        .changelog-item {{ display: flex; gap: 16px; padding: 12px 0; border-bottom: 1px solid #f1f5f9; }}
        .changelog-item:last-child {{ border-bottom: none; }}
        .changelog-version {{ background: #2563eb; color: white; padding: 4px 12px; border-radius: 6px; font-weight: 600; font-size: 0.85em; white-space: nowrap; height: fit-content; }}
        .changelog-date {{ color: #94a3b8; font-size: 0.85em; }}
        .changelog-desc {{ color: #475569; font-size: 0.92em; }}

        /* Footer */
        .footer {{ text-align: center; margin-top: 40px; padding: 30px 40px; background: #f8fafc; border-top: 2px solid #e2e8f0; color: #64748b; }}
        .footer a {{ color: #2563eb; text-decoration: none; font-weight: 500; }}
        .footer .links {{ display: flex; justify-content: center; gap: 20px; margin-top: 12px; flex-wrap: wrap; }}

        @media print {{
            body {{ background: white; padding: 0; }}
            .container {{ box-shadow: none; border-radius: 0; }}
            .header {{ break-after: avoid; }}
            .section {{ break-inside: avoid; }}
        }}
        @media (max-width: 768px) {{
            .content {{ padding: 20px; }}
            .header h1 {{ font-size: 1.8em; }}
            .criteria-grid {{ grid-template-columns: 1fr; }}
            .perf-grid {{ grid-template-columns: repeat(2, 1fr); }}
            .stats-grid {{ grid-template-columns: 1fr; }}
        }}
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <h1>📊 Rapport OLM Benchmark Global</h1>
            <div class="subtitle">Analyse Comparative des Systèmes OCR — olmOCR-Bench {report_data["benchmark_info"]["version"]}</div>
            <div class="date">Généré le {generated}</div>
            <div class="badge">🏛️ {report_data["benchmark_info"]["source"]}</div>
        </div>

        <div class="content">
            <!-- About Section -->
            <div class="section">
                <h2>📋 À propos d'olmOCR-Bench</h2>
                <div class="info-box">
                    <h3>Présentation</h3>
                    <p><strong>Source :</strong> {report_data["benchmark_info"]["source"]}</p>
                    <p><strong>Description :</strong> {report_data["benchmark_info"]["description"]}</p>
                    <p><strong>Documents testés :</strong> {report_data["benchmark_info"]["total_documents"]:,} fichiers PDF</p>
                    <p><strong>Tests unitaires :</strong> {report_data["benchmark_info"]["total_tests"]:,} cas de test</p>
                    <p><strong>Catégories évaluées :</strong> {report_data["benchmark_info"]["test_categories"]} dimensions de performance</p>
                    <p><strong>Licence :</strong> {report_data["benchmark_info"]["license"]}</p>
                    <p><strong>Dernière mise à jour :</strong> {report_data["benchmark_info"]["last_updated"]}</p>
                    <p style="margin-top: 12px;">
                        <strong>Ressources :</strong>
                        <a href="{report_data["benchmark_info"]["github_url"]}" target="_blank">GitHub</a> |
                        <a href="{report_data["benchmark_info"]["huggingface_url"]}" target="_blank">HuggingFace</a> |
                        <a href="{report_data["benchmark_info"]["demo_url"]}" target="_blank">Démo en ligne</a> |
                        <a href="{report_data["benchmark_info"]["arxiv_paper_v1"]}" target="_blank">Paper v1</a> |
                        <a href="{report_data["benchmark_info"]["arxiv_paper_v2"]}" target="_blank">Paper v2</a>
                    </p>
                </div>
            </div>

            <!-- Methodology -->
            <div class="section">
                <h2>🔬 Méthodologie de Scoring</h2>
                <div class="methodology">
                    <h3>Comment les scores sont calculés</h3>
                    <p>{report_data["scoring_methodology"]["overview"]}</p>
                    <p style="margin-top: 8px;">{report_data["scoring_methodology"]["categories_detail"]}</p>
                    <p style="margin-top: 8px;"><strong>Formule :</strong> {report_data["scoring_methodology"]["scoring_formula"]}</p>
                </div>
                <h3>Niveaux de Performance</h3>
                <div class="perf-grid">
                    <div class="perf-item perf-excellent">🟢 Excellent ≥ 90%</div>
                    <div class="perf-item perf-good">🔵 Bon 80-89%</div>
                    <div class="perf-item perf-acceptable">🟡 Acceptable 60-79%</div>
                    <div class="perf-item perf-poor">🔴 Faible &lt; 60%</div>
                </div>
            </div>

            <!-- Criteria -->
            <div class="section">
                <h2>🎯 Critères d'Évaluation</h2>
                <div class="criteria-grid">
"""

    top_by_cat = report_data.get("top_by_category", {})
    for criteria, label in report_data["criteria"]["labels"].items():
        if criteria != "overall":
            desc = report_data["criteria"]["descriptions"][criteria]
            icon = report_data["criteria"]["icons"].get(criteria, "📊")
            leader = top_by_cat.get(criteria)
            leader_html = ""
            if leader:
                leader_html = f'<div class="leader">🏆 Leader : {leader["model_name"]} ({leader["score"]}%)</div>'
            html += f"""
                    <div class="criteria-card">
                        <h4>{icon} {label}</h4>
                        <p>{desc}</p>
                        {leader_html}
                    </div>
"""

    html += """
                </div>
            </div>

            <!-- Benchmark Matrix -->
            <div class="section">
                <h2>🏆 Matrice de Performance — Classement Global</h2>
                <p style="color: #64748b; font-size: 0.9em; margin-bottom: 15px;">
                    (*) Résultats soumis par la communauté, non vérifiés officiellement par AI2.
                </p>
                <div style="overflow-x: auto;">
                <table class="matrix-table">
                    <thead>
                        <tr>
                            <th>Rang</th>
                            <th>Modèle</th>
                            <th>Type</th>
                            <th>ArXiv</th>
                            <th>Math</th>
                            <th>Tables</th>
                            <th>Scans</th>
                            <th>En-têtes</th>
                            <th>Multi-col</th>
                            <th>Texte</th>
                            <th>Base</th>
                            <th>Global</th>
                        </tr>
                    </thead>
                    <tbody>
"""

    for model in report_data["benchmark_matrix"]:
        rank = model["rank"]
        rank_class = f"rank-{rank}" if rank <= 3 else "rank-other"
        scores = model["scores"]
        std_text = f"±{model['std_dev']}" if model.get("std_dev") else ""
        official_mark = "" if model.get("is_official") else "*"
        perf = model.get("performance", get_performance_level(scores["overall"]))

        def score_cls(s):
            if s >= 90:
                return "score-excellent"
            if s >= 80:
                return "score-high"
            if s >= 60:
                return "score-medium"
            return "score-low"

        html += f"""
                        <tr>
                            <td><span class="rank-badge {rank_class}">#{rank}</span></td>
                            <td class="model-name">{model["name"]}{official_mark}</td>
                            <td style="font-size: 0.8em; color: #64748b;">{model.get("model_type", "")}</td>
                            <td class="{score_cls(scores['arxiv'])}">{scores['arxiv']}</td>
                            <td class="{score_cls(scores['old_scans_math'])}">{scores['old_scans_math']}</td>
                            <td class="{score_cls(scores['tables'])}">{scores['tables']}</td>
                            <td class="{score_cls(scores['old_scans'])}">{scores['old_scans']}</td>
                            <td class="{score_cls(scores['headers_footers'])}">{scores['headers_footers']}</td>
                            <td class="{score_cls(scores['multi_column'])}">{scores['multi_column']}</td>
                            <td class="{score_cls(scores['long_tiny_text'])}">{scores['long_tiny_text']}</td>
                            <td class="{score_cls(scores['base'])}">{scores['base']}</td>
                            <td class="{score_cls(scores['overall'])}" style="font-size: 1.1em;">{scores['overall']}{std_text}</td>
                        </tr>
"""

    html += """
                    </tbody>
                </table>
                </div>
            </div>

            <!-- Model Analysis -->
            <div class="section">
                <h2>⭐ Analyse Détaillée par Modèle</h2>
"""

    for model in report_data["benchmark_matrix"][:5]:
        perf = model.get(
            "performance", get_performance_level(model["scores"]["overall"])
        )
        std_text = f"±{model['std_dev']}" if model.get("std_dev") else ""
        html += f"""
                <div class="model-card">
                    <h3>{perf.get("badge","")} {model["name"]} — Rang #{model["rank"]}</h3>
                    <div class="meta">{model.get("model_type", "")} | Score global : {model["scores"]["overall"]}{std_text} | {model.get("note", "")}</div>
"""
        strengths = model.get("strengths", [])
        weaknesses = model.get("weaknesses", [])

        if strengths:
            html += "<p><strong>✅ Points forts :</strong></p><ul>"
            for s in strengths[:4]:
                html += f'<li>{s.get("icon","")} {s["label"]} : <span class="score-excellent">{s["score"]}%</span></li>'
            html += "</ul>"

        if weaknesses:
            html += '<p style="margin-top: 8px;"><strong>⚠️ Points faibles :</strong></p><ul>'
            for w in weaknesses:
                html += f'<li>{w.get("icon","")} {w["label"]} : <span class="score-low">{w["score"]}%</span></li>'
            html += "</ul>"

        html += "</div>"

    html += "</div>"

    info = report_data["benchmark_info"]
    html += f"""
        </div>

        <div class="footer">
            <p><strong>Rapport généré par OCR Intelligence</strong></p>
            <p>Source des données : olmOCR-Bench par {info["source"]}</p>
            <p style="font-size: 0.88em; margin-top: 8px;">
                {info["total_tests"]:,} cas de test sur {info["total_documents"]:,} documents PDF — {info["test_categories"]} catégories d'évaluation
            </p>
            <div class="links">
                <a href="{info["github_url"]}" target="_blank">GitHub</a>
                <a href="{info["huggingface_url"]}" target="_blank">HuggingFace</a>
                <a href="{info["demo_url"]}" target="_blank">Démo</a>
                <a href="{info["arxiv_paper_v1"]}" target="_blank">Paper v1</a>
                <a href="{info["arxiv_paper_v2"]}" target="_blank">Paper v2</a>
            </div>
        </div>
    </div>
</body>
</html>
"""
    return html


def generate_pdf_report(report_data: dict) -> bytes:
    """Generate PDF report from report data using WeasyPrint."""
    try:
        from weasyprint import HTML

        html_content = generate_html_report(report_data)
        pdf_bytes = HTML(string=html_content).write_pdf()
        return pdf_bytes
    except ImportError:
        raise Exception(
            "WeasyPrint n'est pas installé. Installez-le avec: pip install weasyprint"
        )


def generate_docx_report(report_data: dict) -> bytes:
    """Generate Word document report."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

        title = doc.add_heading("Rapport OLM Benchmark Global", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            run.font.color.rgb = RGBColor(30, 58, 138)

        subtitle = doc.add_paragraph(
            "Analyse Comparative des Systèmes OCR — olmOCR-Bench v1.0"
        )
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        generated = datetime.fromisoformat(report_data["generated_at"]).strftime(
            "%d/%m/%Y à %H:%M"
        )
        date_para = doc.add_paragraph(f"Généré le {generated}")
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        doc.add_heading("À propos d'olmOCR-Bench", 1)
        info = report_data["benchmark_info"]
        doc.add_paragraph(f'Source : {info["source"]}')
        doc.add_paragraph(f'Description : {info["description"]}')
        doc.add_paragraph(
            f'Documents testés : {info["total_documents"]:,} fichiers PDF'
        )
        doc.add_paragraph(f'Tests unitaires : {info["total_tests"]:,} cas de test')
        doc.add_paragraph(
            f'Catégories évaluées : {info["test_categories"]} dimensions de performance'
        )
        doc.add_paragraph(f'Licence : {info["license"]}')
        doc.add_paragraph(f'Dernière mise à jour : {info["last_updated"]}')
        doc.add_paragraph(f'GitHub : {info["github_url"]}')
        doc.add_paragraph(f'HuggingFace : {info["huggingface_url"]}')
        doc.add_paragraph(f'Paper v1 : {info["arxiv_paper_v1"]}')
        doc.add_paragraph(f'Paper v2 : {info["arxiv_paper_v2"]}')

        doc.add_page_break()

        doc.add_heading("Méthodologie de Scoring", 1)
        methodology = report_data["scoring_methodology"]
        doc.add_paragraph(methodology["overview"])
        doc.add_paragraph(methodology["categories_detail"])
        doc.add_paragraph(f'Formule : {methodology["scoring_formula"]}')
        doc.add_paragraph()
        doc.add_paragraph("Niveaux de performance :")
        for level, desc in methodology["interpretation"].items():
            doc.add_paragraph(desc, style="List Bullet")

        doc.add_page_break()

        doc.add_heading("Critères d'Évaluation", 1)
        for criteria, label in report_data["criteria"]["labels"].items():
            if criteria != "overall":
                description = report_data["criteria"]["descriptions"][criteria]
                p = doc.add_paragraph(f"{label} : ", style="List Bullet")
                p.add_run(description)

        doc.add_page_break()

        doc.add_heading("Matrice de Performance — Classement Global", 1)
        doc.add_paragraph(
            "(*) Résultats soumis par la communauté, non vérifiés officiellement."
        )

        table = doc.add_table(rows=1, cols=11)
        table.style = "Light Grid Accent 1"

        headers = [
            "Rang",
            "Modèle",
            "ArXiv",
            "Math",
            "Tables",
            "Scans",
            "En-têtes",
            "Multi-col",
            "Texte",
            "Base",
            "Global",
        ]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.size = Pt(9)

        for model in report_data["benchmark_matrix"]:
            row = table.add_row()
            scores = model["scores"]
            std_text = f"±{model['std_dev']}" if model.get("std_dev") else ""
            official_mark = "" if model.get("is_official") else "*"

            row.cells[0].text = f"#{model['rank']}"
            row.cells[1].text = f"{model['name']}{official_mark}"
            row.cells[2].text = str(scores["arxiv"])
            row.cells[3].text = str(scores["old_scans_math"])
            row.cells[4].text = str(scores["tables"])
            row.cells[5].text = str(scores["old_scans"])
            row.cells[6].text = str(scores["headers_footers"])
            row.cells[7].text = str(scores["multi_column"])
            row.cells[8].text = str(scores["long_tiny_text"])
            row.cells[9].text = str(scores["base"])
            row.cells[10].text = f"{scores['overall']}{std_text}"
            for paragraph in row.cells[10].paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        doc.add_page_break()

        doc.add_heading("Analyse Détaillée par Modèle", 1)

        for model in report_data["benchmark_matrix"][:5]:
            std_text = f"±{model['std_dev']}" if model.get("std_dev") else ""
            doc.add_heading(f'{model["name"]} — Rang #{model["rank"]}', 2)
            doc.add_paragraph(
                f'{model.get("model_type", "")} | Score global : {model["scores"]["overall"]}{std_text}'
            )
            doc.add_paragraph(f'Note : {model.get("note", "")}')

            strengths = model.get("strengths", [])
            weaknesses = model.get("weaknesses", [])

            if strengths:
                doc.add_paragraph("Points forts :", style="List Bullet")
                for s in strengths[:4]:
                    doc.add_paragraph(
                        f'{s["label"]} : {s["score"]}%', style="List Bullet 2"
                    )

            if weaknesses:
                doc.add_paragraph("Points faibles :", style="List Bullet")
                for w in weaknesses:
                    doc.add_paragraph(
                        f'{w["label"]} : {w["score"]}%', style="List Bullet 2"
                    )

        doc.add_paragraph()
        footer = doc.add_paragraph(
            "Rapport généré par OCR Intelligence — Source : olmOCR-Bench par Allen Institute for AI"
        )
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()

    except ImportError:
        raise Exception(
            "python-docx n'est pas installé. Installez-le avec: pip install python-docx"
        )


def generate_excel_report(report_data: dict) -> bytes:
    """Generate Excel report with multiple sheets."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

        wb = openpyxl.Workbook()

        header_fill = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )
        header_font = Font(color="FFFFFF", bold=True, size=11)
        title_font = Font(bold=True, size=14, color="1E3A8A")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        green_font = Font(color="059669", bold=True)
        red_font = Font(color="DC2626", bold=True)

        ws = wb.active
        ws.title = "Matrice Benchmark"

        ws.merge_cells("A1:K1")
        ws["A1"] = "📊 Matrice de Performance OLM Benchmark Global"
        ws["A1"].font = Font(bold=True, size=16, color="1E3A8A")
        ws["A1"].alignment = Alignment(horizontal="center")

        generated = datetime.fromisoformat(report_data["generated_at"]).strftime(
            "%d/%m/%Y à %H:%M"
        )
        ws.merge_cells("A2:K2")
        ws["A2"] = (
            f"Généré le {generated} — Source : {report_data['benchmark_info']['source']}"
        )
        ws["A2"].alignment = Alignment(horizontal="center")

        row = 4
        headers = [
            "Rang",
            "Modèle",
            "Type",
            "ArXiv",
            "Math",
            "Tables",
            "Scans",
            "En-têtes",
            "Multi-col",
            "Texte Long",
            "Global",
        ]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        for model in report_data["benchmark_matrix"]:
            row += 1
            scores = model["scores"]
            std_text = f"±{model['std_dev']}" if model.get("std_dev") else ""
            official_mark = "" if model.get("is_official") else "*"

            ws.cell(row=row, column=1, value=model["rank"]).border = border
            ws.cell(
                row=row, column=2, value=f"{model['name']}{official_mark}"
            ).border = border
            ws.cell(row=row, column=3, value=model.get("model_type", "")).border = (
                border
            )

            score_cols = [
                scores["arxiv"],
                scores["old_scans_math"],
                scores["tables"],
                scores["old_scans"],
                scores["headers_footers"],
                scores["multi_column"],
                scores["long_tiny_text"],
            ]
            for i, score in enumerate(score_cols, start=4):
                cell = ws.cell(row=row, column=i, value=score)
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
                if score >= 80:
                    cell.font = green_font
                elif score < 60:
                    cell.font = red_font

            cell = ws.cell(row=row, column=11, value=f"{scores['overall']}{std_text}")
            cell.font = Font(bold=True, size=12, color="1E3A8A")
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

        widths = {
            "A": 8,
            "B": 22,
            "C": 20,
            "D": 10,
            "E": 10,
            "F": 10,
            "G": 10,
            "H": 10,
            "I": 12,
            "J": 12,
            "K": 12,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws2 = wb.create_sheet("Méthodologie")
        ws2.merge_cells("A1:B1")
        ws2["A1"] = "📋 Informations & Méthodologie"
        ws2["A1"].font = title_font

        row = 3
        info = report_data["benchmark_info"]
        meta_items = [
            ("Benchmark", info["name"]),
            ("Version", info["version"]),
            ("Source", info["source"]),
            ("Documents testés", f"{info['total_documents']:,} fichiers PDF"),
            ("Tests unitaires", f"{info['total_tests']:,} cas de test"),
            ("Catégories", f"{info['test_categories']} dimensions"),
            ("Licence", info["license"]),
            ("GitHub", info["github_url"]),
            ("HuggingFace", info["huggingface_url"]),
            ("Paper v1", info["arxiv_paper_v1"]),
            ("Paper v2", info["arxiv_paper_v2"]),
        ]
        for label, value in meta_items:
            ws2.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws2.cell(row=row, column=2, value=value)
            row += 1

        row += 2
        ws2.cell(row=row, column=1, value="Méthodologie de Scoring").font = title_font
        row += 1
        ws2.cell(
            row=row, column=1, value=report_data["scoring_methodology"]["overview"]
        )
        ws2.merge_cells(f"A{row}:B{row}")
        row += 1
        ws2.cell(
            row=row,
            column=1,
            value=report_data["scoring_methodology"]["scoring_formula"],
        )
        ws2.merge_cells(f"A{row}:B{row}")

        ws2.column_dimensions["A"].width = 25
        ws2.column_dimensions["B"].width = 80

        ws3 = wb.create_sheet("Leaders par Catégorie")
        ws3.merge_cells("A1:C1")
        ws3["A1"] = "🏆 Meilleur Modèle par Catégorie"
        ws3["A1"].font = title_font

        row = 3
        headers = ["Catégorie", "Meilleur Modèle", "Score"]
        for col, h in enumerate(headers, start=1):
            cell = ws3.cell(row=row, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border

        for criteria, leader in report_data.get("top_by_category", {}).items():
            row += 1
            ws3.cell(row=row, column=1, value=leader["label"]).border = border
            ws3.cell(row=row, column=2, value=leader["model_name"]).border = border
            cell = ws3.cell(row=row, column=3, value=f"{leader['score']}%")
            cell.border = border
            cell.font = green_font

        ws3.column_dimensions["A"].width = 25
        ws3.column_dimensions["B"].width = 25
        ws3.column_dimensions["C"].width = 12

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    except ImportError:
        raise Exception(
            "openpyxl n'est pas installé. Installez-le avec: pip install openpyxl"
        )


def generate_csv_report(report_data: dict) -> str:
    """Generate CSV report from benchmark data."""
    output = io.StringIO()
    writer = csv.writer(output)

    generated = datetime.fromisoformat(report_data["generated_at"]).strftime(
        "%d/%m/%Y %H:%M"
    )

    writer.writerow(["Rapport OLM Benchmark Global — olmOCR-Bench"])
    writer.writerow([f"Généré le : {generated}"])
    writer.writerow([f"Source : {report_data['benchmark_info']['source']}"])
    writer.writerow([])

    writer.writerow(
        [
            "Rang",
            "Modèle",
            "Type",
            "Officiel",
            "ArXiv",
            "Old Scans Math",
            "Tables",
            "Old Scans",
            "Headers & Footers",
            "Multi Column",
            "Long Tiny Text",
            "Base",
            "Overall",
            "Écart-type",
        ]
    )

    for model in report_data["benchmark_matrix"]:
        scores = model["scores"]
        writer.writerow(
            [
                model["rank"],
                model["name"],
                model.get("model_type", ""),
                "Oui" if model.get("is_official") else "Non",
                scores["arxiv"],
                scores["old_scans_math"],
                scores["tables"],
                scores["old_scans"],
                scores["headers_footers"],
                scores["multi_column"],
                scores["long_tiny_text"],
                scores["base"],
                scores["overall"],
                model.get("std_dev", "N/A"),
            ]
        )

    writer.writerow([])
    writer.writerow(["Leaders par Catégorie"])
    writer.writerow(["Catégorie", "Meilleur Modèle", "Score"])
    for criteria, leader in report_data.get("top_by_category", {}).items():
        writer.writerow([leader["label"], leader["model_name"], f"{leader['score']}%"])

    return output.getvalue()
