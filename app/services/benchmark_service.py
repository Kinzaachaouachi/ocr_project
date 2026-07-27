import io
import json
from datetime import datetime
from typing import Optional

from sqlalchemy import func
from sqlalchemy.orm import Session

from ..models.ocr_history import OCRHistory
from ..models.user import User


def get_user_benchmark_stats(db: Session, user_id: int) -> dict:
    """
    Get aggregated benchmark metrics for a user based on their OCR extractions.
    Returns metrics: precision, execution_time, init_time, robustness, global_score.
    """
    successful = db.query(OCRHistory).filter(
        OCRHistory.user_id == user_id, OCRHistory.status == "success"
    )

    total_extractions = (
        db.query(OCRHistory).filter(OCRHistory.user_id == user_id).count()
    )
    success_count = successful.count()
    failed_count = (
        db.query(OCRHistory)
        .filter(OCRHistory.user_id == user_id, OCRHistory.status == "error")
        .count()
    )

    metrics = successful.with_entities(
        func.avg(OCRHistory.precision_score).label("avg_precision"),
        func.avg(OCRHistory.ocr_time_s).label("avg_exec_time"),
        func.avg(OCRHistory.init_time_s).label("avg_init_time"),
        func.avg(OCRHistory.robustness).label("avg_robustness"),
        func.avg(OCRHistory.global_score).label("avg_global_score"),
    ).first()

    model_stats = (
        successful.with_entities(
            OCRHistory.model_id,
            OCRHistory.model_name,
            func.count(OCRHistory.id).label("count"),
            func.avg(OCRHistory.precision_score).label("avg_precision"),
            func.avg(OCRHistory.ocr_time_s).label("avg_exec_time"),
            func.avg(OCRHistory.init_time_s).label("avg_init_time"),
            func.avg(OCRHistory.robustness).label("avg_robustness"),
            func.avg(OCRHistory.global_score).label("avg_global_score"),
        )
        .group_by(OCRHistory.model_id, OCRHistory.model_name)
        .all()
    )

    file_type_stats = (
        db.query(OCRHistory.file_type, func.count(OCRHistory.id).label("count"))
        .filter(OCRHistory.user_id == user_id)
        .group_by(OCRHistory.file_type)
        .all()
    )

    recent = (
        db.query(OCRHistory)
        .filter(OCRHistory.user_id == user_id)
        .order_by(OCRHistory.processed_at.desc())
        .limit(10)
        .all()
    )

    return {
        "total_extractions": total_extractions,
        "successful_extractions": success_count,
        "failed_extractions": failed_count,
        "success_rate": (
            round((success_count / total_extractions * 100), 2)
            if total_extractions > 0
            else 0
        ),
        "overall_metrics": {
            "precision": round(float(metrics.avg_precision or 0), 2),
            "execution_time": round(float(metrics.avg_exec_time or 0), 2),
            "init_time": round(float(metrics.avg_init_time or 0), 2),
            "robustness": round(float(metrics.avg_robustness or 0), 2),
            "global_score": round(float(metrics.avg_global_score or 0), 2),
        },
        "model_benchmarks": [
            {
                "model_id": s.model_id,
                "model_name": s.model_name,
                "extractions_count": s.count,
                "avg_precision": round(float(s.avg_precision or 0), 2),
                "avg_execution_time": round(float(s.avg_exec_time or 0), 2),
                "avg_init_time": round(float(s.avg_init_time or 0), 2),
                "avg_robustness": round(float(s.avg_robustness or 0), 2),
                "avg_global_score": round(float(s.avg_global_score or 0), 2),
            }
            for s in model_stats
        ],
        "file_type_distribution": {s.file_type: s.count for s in file_type_stats},
        "recent_extractions": [
            {
                "id": r.id,
                "filename": r.filename,
                "model_id": r.model_id,
                "model_name": r.model_name,
                "status": r.status,
                "ocr_time_s": r.ocr_time_s,
                "precision_score": r.precision_score,
                "global_score": r.global_score,
                "processed_at": r.processed_at.isoformat() if r.processed_at else None,
            }
            for r in recent
        ],
    }


def generate_benchmark_csv(db: Session, user_id: int) -> str:
    """Generate CSV content for benchmark export."""
    records = (
        db.query(OCRHistory)
        .filter(OCRHistory.user_id == user_id, OCRHistory.status == "success")
        .order_by(OCRHistory.processed_at.desc())
        .all()
    )

    lines = [
        "Fichier,Date,Modèle,Précision,Temps Exécution,Temps Init,Robustesse,Score Global,Mots,Caractères"
    ]
    for r in records:
        date_str = r.processed_at.strftime("%Y-%m-%d %H:%M") if r.processed_at else ""
        lines.append(
            f'"{r.filename}",{date_str},{r.model_name},'
            f"{r.precision_score or 0},{r.ocr_time_s or 0},{r.init_time_s or 0},"
            f"{r.robustness or 0},{r.global_score or 0},{r.word_count},{r.char_count}"
        )
    return "\n".join(lines)


def generate_benchmark_excel(db: Session, user_id: int) -> bytes:
    """Generate Excel file for benchmark export."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Benchmark OCR"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(
            start_color="1B2A4A", end_color="1B2A4A", fill_type="solid"
        )
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        headers = [
            "Fichier",
            "Date",
            "Modèle OCR",
            "Précision (%)",
            "Temps Exécution (s)",
            "Temps Init (s)",
            "Robustesse",
            "Score Global",
            "Mots",
            "Caractères",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        records = (
            db.query(OCRHistory)
            .filter(OCRHistory.user_id == user_id, OCRHistory.status == "success")
            .order_by(OCRHistory.processed_at.desc())
            .all()
        )

        for row_idx, r in enumerate(records, 2):
            data = [
                r.filename,
                r.processed_at.strftime("%Y-%m-%d %H:%M") if r.processed_at else "",
                r.model_name,
                r.precision_score or 0,
                r.ocr_time_s or 0,
                r.init_time_s or 0,
                r.robustness or 0,
                r.global_score or 0,
                r.word_count,
                r.char_count,
            ]
            for col, val in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")

        for col in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    except ImportError:

        csv_content = generate_benchmark_csv(db, user_id)
        return csv_content.encode("utf-8-sig")
